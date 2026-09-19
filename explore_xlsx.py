import openpyxl

PATH = "C:/Users/HP/Documents/Github/family-financial-os/Documents/DEUDAS PROYECCION 2026.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)

print("=== SHEETS ===")
for name in wb.sheetnames:
    print(name)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== SHEET: {sheet_name} (rows={ws.max_row}, cols={ws.max_column}) ===")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50), values_only=False):
        vals = []
        for cell in row:
            v = cell.value
            if v is not None:
                vals.append(f"{cell.coordinate}={v}")
        if vals:
            print("  " + " | ".join(vals))
